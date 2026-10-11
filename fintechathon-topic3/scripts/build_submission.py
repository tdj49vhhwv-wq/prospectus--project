"""Create five Chinese figures and an eight-page-body competition report.

PDF generation uses XeLaTeX; DOCX and Markdown remain editable. Business names,
review text and reviewer identifiers are never included in the report.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.font_manager import FontProperties
import numpy as np
import pandas as pd
import requests

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "outputs"
REPORT = ROOT / "report"
FONT_BASE = "https://raw.githubusercontent.com/Yixf-Self/fandol-fonts/b93300821373d3092e378e57ae22ecb8e9082c62/"
FONTS = {
    "FandolSong-Regular.otf": ("FandolSong-Regular/FandolSong-Regular.otf", "5dc10d4d76ec3bc5fcc6065d80088c38eee0cbf851bee5b5cad1334a796a0d59"),
    "FandolSong-Bold.otf": ("FandolSong-Bold/FandolSong-Bold.otf", "ba71e3247d04ffcd547f5826c7eb747da4f70cfe047c166283f5f31c402817eb"),
    "COPYING": ("COPYING", "853b586f0d520493390e571431afaf36a5fbb27dcfd239338a7ee9b0505cb004"),
}


def ensure_fonts():
    folder = REPORT / "fonts"
    folder.mkdir(parents=True, exist_ok=True)
    for name, (remote, digest) in FONTS.items():
        file = folder / name
        if not file.exists():
            response = requests.get(FONT_BASE + remote, timeout=60)
            response.raise_for_status()
            if hashlib.sha256(response.content).hexdigest() != digest:
                raise RuntimeError("Font checksum mismatch: " + name)
            file.write_bytes(response.content)
        assert hashlib.sha256(file.read_bytes()).hexdigest() == digest, name
    return FontProperties(fname=str(folder / "FandolSong-Regular.otf"))


def figures(font):
    folder = REPORT / "figures"
    folder.mkdir(parents=True, exist_ok=True)
    metrics = pd.read_csv(OUT / "model_metrics.csv", index_col="model")
    summary = json.loads((OUT / "analysis_summary.json").read_text())
    audit = json.loads((OUT / "data_audit.json").read_text())
    desc = pd.read_csv(OUT / "descriptive_stats.csv").set_index("feature")
    budgets = pd.read_csv(OUT / "review_budget.csv")
    sensitivity = pd.read_csv(OUT / "sensitivity_checks.csv")
    colors = {"open": "#296A88", "closed": "#B9573C", "primary": "#724D99", "tree": "#287E62"}
    plt.rcParams.update({"font.size": 10, "axes.unicode_minus": False, "axes.spines.top": False, "axes.spines.right": False, "figure.dpi": 160, "savefig.dpi": 300})

    def finish(fig, filename):
        for ax in fig.axes:
            for text in [ax.title, ax.xaxis.label, ax.yaxis.label, *ax.get_xticklabels(), *ax.get_yticklabels(), *ax.texts]:
                text.set_fontproperties(font)
            legend = ax.get_legend()
            if legend:
                for text in legend.get_texts():
                    text.set_fontproperties(font)
        for text in fig.texts:
            text.set_fontproperties(font)
        fig.tight_layout(pad=0.8)
        fig.savefig(folder / filename, facecolor="white", bbox_inches="tight")
        fig.savefig(folder / filename.replace(".png", ".pdf"), facecolor="white", bbox_inches="tight")
        plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(7.05, 2.5), gridspec_kw={"width_ratios": [1.6, 1]})
    values = [3164, 1361, 900, 851, 674]
    labels = ["区域全部商户", "餐饮商户", "总评论≥10条", "历史跨度≥365天", "近期评论≥3条"]
    axes[0].barh(np.arange(5), values, color=["#AAC0CB", "#839EA9", "#63879B", "#497A91", "#296A88"], height=.6)
    axes[0].set_yticks(np.arange(5), labels); axes[0].invert_yaxis(); axes[0].set_xlim(0, 3600)
    for i, value in enumerate(values):
        axes[0].text(value + 45, i, f"{value:,}", va="center", fontsize=10)
    axes[0].set_xlabel("逐步筛选后商户数")
    axes[1].bar([0, 1], [551, 123], color=[colors["open"], colors["closed"]], width=.6)
    axes[1].set_xticks([0, 1], ["营业", "停止营业"]); axes[1].set_ylim(0, 640)
    for i, value in enumerate([551, 123]):
        axes[1].text(i, value + 15, str(value), ha="center")
    axes[1].set_ylabel("最终样本商户数"); axes[1].text(.5, 600, "停止营业占比18.25%", ha="center", fontsize=9)
    finish(fig, "figure01_sample.png")

    fig, axes = plt.subplots(1, 4, figsize=(7.05, 2.3))
    columns = ["rating_mean_360", "rating_change_180", "recent_reviews_180", "price_negative_share"]
    titles = ["近360天均分", "近期评分变化", "近期评论数", "低分价格主题占比"]
    for ax, column, title in zip(axes, columns, titles):
        row = desc.loc[column]
        values = [row.open_mean, row.closed_mean]
        if column.endswith("share"):
            values = [value * 100 for value in values]
        ax.bar([0, 1], values, color=[colors["open"], colors["closed"]], width=.6)
        ax.set_xticks([0, 1], ["营业", "停业"]); ax.set_title(title, fontsize=9)
        ax.axhline(0, color="#555555", lw=.6)
        for i, value in enumerate(values):
            label = f"{value:.2f}%" if column.endswith("share") else (f"{value:.3f}" if column == "rating_change_180" else f"{value:.2f}")
            ax.annotate(label, (i, value), xytext=(0, 3 if value >= 0 else -11), textcoords="offset points", ha="center", fontsize=9)
        if column == "rating_change_180":
            ax.set_ylim(-.35, .13)
        else:
            ax.set_ylim(0, max(values) * 1.35)
        ax.tick_params(labelsize=9)
    finish(fig, "figure02_descriptive.png")

    fig, axes = plt.subplots(1, 2, figsize=(7.05, 3.0), sharey=True)
    names = list(metrics.index)
    labels = [name.replace("_", "·") + (" 主模型" if name == "LR_QAT" else "") for name in names]
    ci = summary["confidence_intervals"]["per_model"]
    for ax, measure, title in zip(axes, ["roc_auc", "pr_auc_ap"], ["测试ROC-AUC", "测试平均精确率AP"]):
        for i, name in enumerate(names):
            value = metrics.loc[name, measure]
            lower, upper = ci[name][measure]
            color = colors["primary"] if name == "LR_QAT" else (colors["tree"] if name.startswith("HGB") else colors["open"])
            ax.errorbar(value, i, xerr=[[value - lower], [upper - value]], fmt="o", color=color, capsize=3, markersize=5)
            ax.text(.97, i, f"{value:.3f}", va="center", ha="right", fontsize=9)
        reference = .5 if measure == "roc_auc" else summary["test_prevalence"]
        ax.axvline(reference, color="#777777", ls="--", lw=1)
        ax.axhspan(1.6, 2.4, color=colors["primary"], alpha=.08)
        ax.set_xlim(0, 1); ax.set_title(title); ax.set_xticks([0, .25, .5, .75, 1])
        ax.set_yticks(np.arange(len(names)), labels)
        ax.grid(axis="x", alpha=.15)
    axes[0].invert_yaxis()
    finish(fig, "figure03_models.png")

    fig, axes = plt.subplots(1, 2, figsize=(7.05, 2.65), gridspec_kw={"width_ratios": [1.3, 1]})
    strategies = [("random_expected", "随机期望", "#888888", "--"), ("low_rating", "低评分筛选", "#B88637", "-"), ("LR_QAT", "增强逻辑回归（主模型）", colors["primary"], "-"), ("HGB_QAT", "增强树模型（探索比较）", colors["tree"], "-")]
    for name, label, color, style in strategies:
        part = budgets.loc[budgets.strategy == name].sort_values("budget_fraction")
        axes[0].plot(part.reviewed / 135 * 100, part.recall * 100, color=color, ls=style, marker="o", label=label, lw=1.6)
    axes[0].set_xlabel("实际复核比例（%）"); axes[0].set_ylabel("停业商户捕获率（%）")
    axes[0].set_ylim(0, 55); axes[0].set_xlim(8, 33); axes[0].legend(fontsize=8, loc="upper left"); axes[0].grid(alpha=.15)
    names = ["随机期望", "低评分", "主模型", "增强树模型"]
    hit = [25 * 14 / 135, 3, 3, 7]
    axes[1].barh(np.arange(4), hit, color=["#AAAAAA", "#B88637", colors["primary"], colors["tree"]], height=.58)
    for i, value in enumerate(hit):
        axes[1].text(value + .15, i, f"{value:.2f}" if i == 0 else str(int(value)), va="center", fontsize=9)
    axes[1].set_yticks(np.arange(4), names); axes[1].invert_yaxis(); axes[1].set_xlim(0, 9)
    axes[1].set_xlabel("复核14家时命中数"); axes[1].set_title("10.37%实际复核预算", fontsize=10)
    finish(fig, "figure04_budget.png")

    fig, axes = plt.subplots(1, 2, figsize=(7.05, 2.5), gridspec_kw={"width_ratios": [1, 1.15]})
    years = audit["anchor_year_by_status"]
    strata = [("2014及以前", lambda y: y <= 2014), ("2015—2016", lambda y: 2015 <= y <= 2016), ("2017", lambda y: y == 2017)]
    left = np.zeros(2)
    for (label, predicate), color in zip(strata, ["#BAC6D0", "#D9B26A", "#4D879B"]):
        amounts = np.array([sum(value for year, value in years[status].items() if predicate(int(year))) for status in ["0", "1"]])
        proportions = amounts / np.array([551, 123]) * 100
        axes[0].barh([0, 1], proportions, left=left, label=label, color=color, height=.5)
        for i, proportion in enumerate(proportions):
            if proportion >= 10:
                axes[0].text(left[i] + proportion / 2, i, f"{proportion:.1f}%", ha="center", va="center", fontsize=9)
        left += proportions
    axes[0].set_yticks([0, 1], ["营业", "停业"]); axes[0].invert_yaxis(); axes[0].set_xlim(0, 100)
    axes[0].set_xlabel("最后评论年份构成（%）"); axes[0].legend(fontsize=8, loc="lower center", bbox_to_anchor=(.5, -.58), ncol=3)
    order = ["all_test", "at_least_20_lifetime_reviews", "exact_Cleveland_city", "last_review_2015_or_later"]
    names = ["全部测试（135家）", "总评论≥20（114家）", "核心城区（126家）", "评论年份≥2015（125家）"]
    subset = sensitivity.set_index("subset").loc[order]
    axes[1].scatter(subset.roc_auc, np.arange(4), color=colors["primary"], s=30)
    for i, value in enumerate(subset.roc_auc):
        axes[1].text(value + .005, i, f"{value:.3f}", va="center", fontsize=9)
    axes[1].axvline(.5, color="#888888", ls="--"); axes[1].set_xlim(.46, .68)
    axes[1].set_yticks(np.arange(4), names); axes[1].invert_yaxis(); axes[1].set_xlabel("主模型测试ROC-AUC")
    finish(fig, "figure05_bias.png")


def paragraph(text):
    return {"type": "paragraph", "text": text.replace("−", "-")}


def heading(text):
    return {"type": "heading", "text": text}


def table(title, headers, rows, widths=None):
    return {"type": "table", "title": title, "headers": headers, "rows": rows, "widths": widths}


def figure(name, caption, width=1):
    return {"type": "figure", "file": name, "caption": caption, "width": width}


def content():
    title = "基于公开口碑信息的餐饮商户经营风险画像及验证"
    subtitle = "——营业状态识别的可行性研究"
    pages = [
        [heading("摘要"),
         paragraph("针对小微餐饮主体财务信息稀缺的问题，本文研究公开口碑信息能否补充经营风险画像。基于某生活服务平台的海外历史区域子集，筛选674家餐饮商户，其中123家记录为停止营业；围绕口碑质量、评论活跃度和低分评论主题构建12项指标，比较逻辑回归与梯度提升树的六种输入组合，并进行独立测试和人工复核预算模拟。按验证集确定的增强逻辑回归，测试ROC-AUC为0.577、平均精确率AP为0.220；相对评分质量基线的AP增量为0.009，95%配对Bootstrap区间包含0。预设增强树模型AP为0.319，在复核14家测试商户时命中7家停止营业商户，相对随机名单的点估计提升2.70倍，但区间仍包含无提升的可能。公开信息可形成低成本调查线索，尚未证明稳定的风险识别增益。由于缺少关闭日期、真实流水和信用表现，结果仅支持历史营业状态识别，不能解释为提前停业预测或信用违约预测。"),
         paragraph("关键词：公开数据；餐饮商户；经营风险画像；营业状态；人工复核"),
         heading("一、业务背景与研究问题"),
         paragraph("小微经营主体现金流缓冲有限，传统财务风控又依赖标准化报表和信用记录。生活服务平台的公开评价覆盖消费体验、商户活跃度及经营环境，为信息稀疏主体提供了另一条观察路径。但公开评分与经营风险的关系可能受到平台选择、商户规模和时间结构影响，必须用独立样本验证，不能把低评分直接等同于高信用风险。"),
         paragraph("本文提出三个问题：第一，能否将公开评论转化为可解释的经营画像指标？第二，相较于只看评分质量，评论活跃度和低分主题是否提供增量识别价值？第三，在人工核查预算有限时，模型排序是否比随机抽查或低评分筛选更有效？研究以餐饮业为代表，以公开营业状态作为可得的替代验证标签。"),
         paragraph("赛题强调在经营流水基础上引入公开信息。本研究未获得真实流水，因此检验的是公开信息的独立识别价值，尚未验证其对流水模型的增益。样本也缺少企业规模和贷款表现，餐饮商户仅作为小微主体的研究代理。停业可能来自搬迁或主动退出等原因，不等于财务困境。")],
        [heading("二、数据介绍与样本清洗"),
         paragraph("原始来源为某生活服务平台的学术开放数据[1]，实际使用其社区发布的区域子集[2]。子集以城市名称包含“leveland”的规则筛选，覆盖Cleveland及邻近区域。本研究固定数据来源提交并校验两个输入文件，独立计算指标和模型，没有复用来源项目的模型结果。商户表含类别、城市和营业状态；评论表含商户ID、日期、评分和文本。"),
         paragraph("原始数据含3,164家商户、91,057条评论。评论日期为2005年4月15日至2017年7月26日，营业状态的采集日期未说明。1,361家餐饮商户中，318家标记停止营业；其评论共64,693条。按照评价数量、历史跨度及近期可观测性逐步筛选，最终纳入674家，其中营业551家、停止营业123家，停止营业占比18.25%。"),
         figure("figure01_sample.png", "图1 样本逐步筛选及最终营业状态分布"),
         table("表1 数据质量与样本纳入规则", ["核对项目", "结果或处理规则"], [
             ["商户与评论关联", "商户ID无重复；评论ID无重复；评论均匹配商户"],
             ["日期与评分", "无无效日期；评分均处于1—5分"],
             ["纳入门槛", "总评论≥10条；历史跨度≥365天；近180天评论≥3条"],
             ["计数一致性", "64家商户表总评论数与实际评论数不一致，原因未确认"],
             ["计数处理", "所有计数从实际评论重新计算，不使用商户表总数"],
         ], [.27, .73]),
         paragraph("最终样本包含核心城区商户635家、邻近城区39家。筛选排除了新店、评价较少及长期无评论的商户，改善指标可观测性，也形成样本选择偏差。全部报告图表均为匿名聚合结果，不展示商户名称、评论者信息或原始评论。")],
        [heading("三、指标体系与描述性分析"),
         paragraph("以各商户最后一条评论日期作为历史窗口锚点：近360天为距锚点0—359天，近期与前期分别为0—179天和180—359天。窗口在不同商户之间不对应统一日历时间，因而本设计属于历史横截面比较。基于经营机制构建三类、共12项指标，见表2。"),
         table("表2 三维经营风险画像指标体系", ["维度", "指标（每维4项）", "业务含义"], [
             ["口碑质量Q", "近360天均分、低分占比、评分标准差；近期较前期均分变化", "衡量消费者评价水平、分歧及变化"],
             ["评论活跃度A", "近期/前期评论数；数量对数变化；六个30天窗口数量变异系数", "衡量平台消费关注度及波动"],
             ["低分主题T", "近期低分且提及服务、卫生、价格、食物问题的评论占比", "识别可供调查核实的经营体验线索"],
         ], [.18, .48, .34]),
         paragraph("数量变化定义为log(1+近期数量)−log(1+前期数量)。低分指评分不超过2分；各主题以固定词典匹配，分母为近期全部评论数。40家商户前期窗口无评论，其均分变化记为缺失，在拟合数据上作中位数填补并增加缺失指示。评论数不是客流或营业额，词典主题也未经独立人工标注验证。"),
         figure("figure02_descriptive.png", "图2 营业与停止营业商户的主要指标均值（全样本探索性描述）"),
         paragraph("停止营业商户近360天均分为3.43，营业商户为3.66；近期评分变化均值分别为−0.265与0.002。近期评论数均值分别为7.63与11.21，但两组中位数均为7，说明分布偏态。价格主题占比在停业组较高，卫生主题则未呈现预期方向。组间差异以Mann–Whitney检验并作12项Holm校正，仅作探索性描述，不用于因果推断。")],
        [heading("四、模型分析与验证设计"),
         paragraph("以商户表记录的停止营业为正类，营业为负类。按商户独立、标签分层，一次固定划分训练、验证、测试集合，随机种子为20261011。同一商户不会跨集合，测试集保留到参数与主分析方案确定后使用；划分比例为60%/20%/20%，见表3。"),
         table("表3 样本划分与标签构成", ["集合", "商户数", "停止营业数", "占比"], [
             ["训练集", "404", "73", "18.07%"],
             ["验证集", "135", "25", "18.52%"],
             ["测试集", "135", "25", "18.52%"],
         ], [.31, .23, .23, .23]),
         paragraph("采用逻辑回归（LR）作解释性基线，以及直方图梯度提升树（HGB）检验非线性关系。每种算法均比较Q、Q+A、Q+A+T三组输入，共六个预设模型，以分离评分、活跃度与评论主题的增量。没有将当前营业状态文字、关闭词汇、最后评论年份或距数据末日的无评论天数放入模型。"),
         paragraph("逻辑回归在C=0.1、1、10中选择正则强度；树模型比较最大叶节点7、15及L2系数1、5，固定150次迭代、学习率0.05、每叶最少20个样本。每种算法与指标组仅用验证集平均精确率AP选参数，再合并训练和验证集重新拟合。逻辑回归标准化、所有缺失填补及缺失指示均只在拟合数据估计。"),
         paragraph("质量基线在LR·Q与HGB·Q中按验证集AP选择，增强方案在LR·QAT与HGB·QAT中同样选择。由此主分析基线为LR·Q，主分析增强模型为LR·QAT。增强树模型虽然测试表现较好，但没有据此回头改选为主模型。全部预设模型均保留，并明确区分主分析与探索比较。"),
         paragraph("验证使用ROC-AUC、平均精确率AP和有限预算下的复核命中率、捕获率及提升倍数。AP采用非插值平均精确率，不混用梯形积分PR面积；测试停止营业比例18.52%作为随机排序AP的总体参考。对测试商户作1,500次配对Bootstrap，报告百分位95%区间；该区间不涵盖训练重拟合、模型选择或外部地区的不确定性。")],
        [heading("五、模型结果与增量价值"),
         paragraph("六个模型的独立测试点估计和Bootstrap区间见图3。质量基线LR·Q测试ROC-AUC为0.545、AP为0.211；主分析增强LR·QAT分别为0.577和0.220。加入活跃度与主题后的改善较小，不能据此声称公开信息已显著提升风险识别。"),
         figure("figure03_models.png", "图3 六个预设模型的测试表现及95%Bootstrap区间；虚线为随机参考"),
         table("表4 预设模型的验证与测试结果", ["模型", "验证AP", "测试ROC-AUC", "测试AP"], [
             ["LR·Q", "0.301", "0.545", "0.211"],
             ["LR·QA", "0.320", "0.552", "0.212"],
             ["LR·QAT（主模型）", "0.344", "0.577", "0.220"],
             ["HGB·Q", "0.259", "0.585", "0.253"],
             ["HGB·QA", "0.323", "0.614", "0.299"],
             ["HGB·QAT", "0.304", "0.631", "0.319"],
         ], [.37, .19, .24, .20]),
         paragraph("主模型相对质量基线的AP增量为0.0091，配对95%区间为[−0.055，0.064]。增强树模型相对其质量基线的AP增量为0.0658，区间为[−0.038，0.185]。两者均包含0，无法排除无增益的情况。树模型中QA相对Q、QAT相对QA的AP点估计均有所增加，可作为扩大样本后进一步验证的线索。"),
         paragraph("本次测试仅含25家停止营业商户，区间较宽；多模型比较未作家族显著性校正，结果应按探索性证据理解。模型输出仅表示历史样本的状态排序，不视为校准的未来停业概率或贷款违约概率，也不据测试最高分直接决定部署方案。")],
        [heading("六、应用验证：有限预算下的人工复核"),
         paragraph("设定金融机构每期只能核实约10%的商户。对135家测试商户向上取整选14家，实际预算比例为10.37%，对比随机抽查、最低评分筛选、主分析方案及预设增强树模型。停止营业标签仅用于事后核对名单，未用于选择名单或调整模型参数。"),
         figure("figure04_budget.png", "图4 不同复核预算下的停业商户捕获率及约10%预算的命中数"),
         table("表5 复核14家测试商户的模拟效果", ["策略", "命中数", "名单命中率", "停业捕获率", "提升倍数"], [
             ["随机抽查（期望）", "2.59", "18.52%", "10.37%", "1.00"],
             ["低评分筛选", "3", "21.43%", "12.00%", "1.16"],
             ["质量基线LR·Q", "3", "21.43%", "12.00%", "1.16"],
             ["主模型LR·QAT", "3", "21.43%", "12.00%", "1.16"],
             ["增强树模型HGB·QAT", "7", "50.00%", "28.00%", "2.70"],
         ], [.35, .14, .19, .19, .13]),
         paragraph("随机名单2.59家为理论期望，不是某次抽查实际命中了小数家。主模型与低评分名单均命中3家，未表现出额外复核效率。增强树模型命中7家，提升倍数点估计为2.70，但95%区间为[0.96，3.75]，包含1；模型相对随机的优势仍需独立复验。此树模型属于预设探索比较，不能因测试表现高而追认为已定型的部署模型。"),
         paragraph("增强树模型在约10%预算下仍漏掉72%的停止营业商户，且名单中一半商户仍在营业。因此这类输出更适合安排低成本人工核实，而非直接拒贷。该模拟也没有衡量真实违约、回收损失、核查成本或干预效果，不能将点估计转化为确定的商业收益。")],
        [heading("七、偏差检查、敏感性与研究限制"),
         paragraph("营业与停止营业组的最后评论年份存在明显差异：551家营业商户中528家最后评论在2017年，123家停止营业商户中只有6家在2017年。图5展示这一结构及固定主模型的子样本敏感性结果。日期没有作为模型输入，但窗口年代差异仍可能影响评论量与消费行为。"),
         figure("figure05_bias.png", "图5 日期结构差异及固定主模型的子样本敏感性核对"),
         paragraph("将测试样本限定为最后评论年份2015年及之后时，125家商户中15家标记停业，主模型ROC-AUC为0.525、AP为0.141，接近随机表现。至少20条总评论的114家子样本ROC-AUC为0.558。上述核对只筛选原测试预测，没有为提高分数重新调参；未证明主分析增益稳定。"),
         paragraph("标签边界：没有关闭日期、统一状态采集日期、贷款逾期或真实流水，无法验证评论在风险事件之前可用。随机商户划分只检验同类历史记录的状态识别，不代表跨时间预警能力。停业也不能直接作为违约标签。"),
         paragraph("样本边界：平台消费者并非全部客群，评价数不等于营业额。新店、低评价量和长期无评论主体被排除；企业规模、连锁关系和评论真实性尚未核验。词典主题只反映低分评论中的话题提及，不是精确投诉归因。"),
         paragraph("外推与使用边界：评论截至2017年，海外区域样本不能证明2026年深圳小微企业适用性。当前官方数据协议及该镜像的再分发授权尚待核实，2015年历史协议[3]不能代替当前许可。本报告不披露原始评论及具名商户，正式提交前应确认教育比赛使用范围。")],
        [heading("八、产品思路与商业应用"),
         paragraph("建议将产品定位为“小微商户经营线索助手”：把公开信息转化为调查优先级、触发原因和数据可信度，服务贷前信息补充与贷后人工核查。现阶段不输出“违约概率”，也不以平台评分直接决定拒贷。公开数据模块应与真实流水模块并列，保留各自的信息来源和不确定性。"),
         table("表6 产品模块、输出与业务动作", ["模块", "输出", "建议动作"], [
             ["数据可用性", "评论数量、覆盖跨度、更新情况", "不足时提示“信息不足”，补充资料"],
             ["经营画像", "评分变化、活跃度、低分主题", "核实具体经营问题，查看实际流水"],
             ["调查排序", "预算对应的待复核名单及原因", "人工确认营业情况，记录误报与漏报"],
             ["反馈与验证", "真实调查结果、成本与后续表现", "定期复验，检验增量及商业价值"],
         ], [.22, .35, .43]),
         paragraph("业务落地分三步。第一步先核实经营主体规模、统一观测时间和数据使用范围；第二步在真实流水基线上增加公开指标，采用时间外及地区外测试，比较增量和复核效率；第三步以实际核查成本、风险损失和干预结果评价收益，而不是只看AUC。预算比例由业务能力确定，不能用本次测试反复挑选最有利的阈值。"),
         heading("九、总结结论"),
         paragraph("第一，公开评论可以构建三维12项可解释的经营画像，便于在信息稀疏条件下形成调查线索。第二，本次主分析方案没有显示稳定的增量识别价值；增强树模型的探索性点估计较好，但模型差异和复核提升区间仍包含无改善，不能据此证明已具备金融风控能力。第三，日期混杂、标签缺失和历史地域限制决定了研究更适合可行性检验，尚不能支持自动授信或前瞻预警。"),
         paragraph("公开数据的合理价值，是让授信调查更有针对性，并以可验证的方式补充真实经营信息。下一阶段应优先取得统一时点的真实流水与未来经营/信用标签，完成独立复验后再评估商业部署。本文同时报告积极线索与不确定性，避免将替代标签、单次高分或历史状态痕迹误当成信用风险预测能力。")],
    ]
    references = [heading("参考文献与数据来源"),
        paragraph("[1] Yelp. Yelp Open Dataset[EB/OL]. https://www.yelp.com/dataset. 官方当前使用条款在本次研究中尚未核实。"),
        paragraph("[2] Cleveland区域公开子集及筛选说明[DS/OL]. https://github.com/ahegel/yelp-dataset. 本研究固定提交e5b82be3c909f9bb45dfd03e164ae0b7897832bd；评论截至2017年7月26日，营业状态采集时间未说明。"),
        paragraph("[3] Yelp. Dataset Challenge Academic Dataset Terms of Use[EB/OL]. 更新日期2015-03-23. 历史协议归档来源：https://github.com/vc1492a/Yelp-Challenge-Dataset. 固定提交0402872665e57c253ba7dcbb71b10842bc1ef08a；文件位于该仓库Raw Data目录，题名为Dataset Challenge Academic Dataset Agreement。该历史归档不证明当前授权。"),
        paragraph("[4] scikit-learn. Model evaluation: quantifying the quality of predictions; LogisticRegression; HistGradientBoostingClassifier[EB/OL]. https://scikit-learn.org/stable/. 本研究AP采用average_precision_score的非插值平均精确率定义。"),
        paragraph("[5] Efron B, Tibshirani R J. An Introduction to the Bootstrap[M]. New York: Chapman & Hall, 1993."),
        paragraph("[6] Holm S. A Simple Sequentially Rejective Multiple Test Procedure[J]. Scandinavian Journal of Statistics, 1979, 6(2): 65–70."),
    ]
    return title, subtitle, pages, references


def escape(text):
    replacements = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
    return "".join(replacements.get(character, character) for character in text)


def tex_block(block):
    kind = block["type"]
    if kind == "heading":
        return r"\section*{" + escape(block["text"]) + "}\n"
    if kind == "paragraph":
        parts = re.split(r"(https?://[A-Za-z0-9./_-]+)", block["text"])
        return "".join(r"\url{" + part.rstrip(".") + "}" + ("." if part.endswith(".") else "") if part.startswith("https://") else escape(part) for part in parts) + "\n\n"
    if kind == "figure":
        return "\\begin{figure}[H]\n\\centering\n\\includegraphics[width=" + str(block["width"]) + r"\linewidth]{figures/" + block["file"].replace(".png", ".pdf") + "}\n\\caption{" + escape(block["caption"]) + "}\n\\end{figure}\n"
    if kind == "table":
        widths = block["widths"] or [1 / len(block["headers"])] * len(block["headers"])
        specification = "".join(r">{\raggedright\arraybackslash}p{" + f"{width:.3f}" + r"\tablewidth}" for width in widths)
        lines = ["\\begin{table}[H]", "\\centering", "\\caption{" + escape(block["title"]) + "}", "\\setlength{\\tablewidth}{\\linewidth}", f"\\addtolength{{\\tablewidth}}{{-{2 * len(widths)}\\tabcolsep}}", "\\begin{tabular}{" + specification + "}", r"\toprule"]
        lines.append(" & ".join(r"\textbf{" + escape(str(x)) + "}" for x in block["headers"]) + r" \\")
        lines.append(r"\midrule")
        lines.extend(" & ".join(escape(str(x)) for x in row) + r" \\" for row in block["rows"])
        lines.extend([r"\bottomrule", r"\end{tabular}", r"\end{table}"])
        return "\n".join(lines) + "\n"
    raise ValueError(kind)


def write_tex(title, subtitle, pages, references, team_name):
    preamble = r"""\documentclass[a4paper,10pt]{article}
\usepackage{fontspec}
\setmainfont[Path=fonts/,BoldFont=FandolSong-Bold.otf,ItalicFont=FandolSong-Regular.otf]{FandolSong-Regular.otf}
\XeTeXlinebreaklocale "zh"
\XeTeXlinebreakskip=0pt plus 1pt
\usepackage[left=24mm,right=24mm,top=22mm,bottom=22mm]{geometry}
\usepackage{graphicx,float,booktabs,array,caption,titlesec,fancyhdr,indentfirst}
\usepackage[hidelinks,unicode]{hyperref}
\hypersetup{pdfauthor={},pdftitle={公开口碑信息对餐饮商户经营风险画像的验证},pdfsubject={赛题三，队名待填写的报告稿}}
\renewcommand{\normalsize}{\fontsize{10.5bp}{15.75bp}\selectfont}
\normalsize
\setlength{\parindent}{21bp}
\setlength{\parskip}{0pt}
\setlength{\tabcolsep}{4pt}
\newlength{\tablewidth}
\renewcommand{\arraystretch}{1.05}
\setlength{\textfloatsep}{6pt}
\setlength{\intextsep}{6pt}
\captionsetup{labelformat=empty,font=normalsize,skip=3pt}
\titleformat{\section}{\normalfont\fontsize{12pt}{18pt}\selectfont\bfseries}{}{0pt}{}
\titlespacing*{\section}{0pt}{6pt}{4pt}
\pagestyle{fancy}
\fancyhf{}
\fancyfoot[C]{\fontsize{9pt}{12pt}\selectfont\thepage}
\renewcommand{\headrulewidth}{0pt}
\setlength{\headheight}{0pt}
\setlength{\footskip}{12mm}
\emergencystretch=1em
\begin{document}
\begin{titlepage}
\centering
\vspace*{30mm}
{\fontsize{15pt}{23pt}\selectfont 2026 FinTechathon 深圳国际金融科技大赛\par}
\vspace{5mm}
{\fontsize{12pt}{18pt}\selectfont 数据分析赛道 · 赛题三\par}
\vspace{26mm}
"""
    cover = "{\\fontsize{20pt}{30pt}\\selectfont\\bfseries " + escape(title) + "\\par}\n\\vspace{5mm}\n{\\fontsize{13pt}{21pt}\\selectfont " + escape(subtitle) + "\\par}\n\\vspace{30mm}\n{\\fontsize{14pt}{21pt}\\selectfont 小组名称：" + escape(team_name) + "\\par}\n\\end{titlepage}\n\\setcounter{page}{1}\n\\normalsize\n"
    body = "\n\\newpage\n".join("".join(tex_block(block) for block in page) for page in pages)
    tail = "\n\\newpage\n\\fancyfoot[C]{\\fontsize{9pt}{12pt}\\selectfont 参考文献}\n" + "".join(tex_block(block) for block in references) + "\n\\end{document}\n"
    (REPORT / "report.tex").write_text(preamble + cover + body + tail, encoding="utf-8")


def markdown(title, subtitle, pages, references, team_name):
    lines = ["# " + title, subtitle, "", "小组名称：" + team_name, "", "<!-- 封面后正文共8页；参考文献另页；PDF页数已经单独核验。 -->", ""]
    for page_number, blocks in enumerate([*pages, references], 1):
        lines.append(f"<!-- {'正文第'+str(page_number)+'页' if page_number <= 8 else '参考文献'} -->")
        for block in blocks:
            if block["type"] == "heading":
                lines.extend(["## " + block["text"], ""])
            elif block["type"] == "paragraph":
                lines.extend([block["text"], ""])
            elif block["type"] == "figure":
                lines.extend([f"![{block['caption']}](figures/{block['file']})", "", block["caption"], ""])
            else:
                lines.extend([block["title"], "", "| " + " | ".join(block["headers"]) + " |", "| " + " | ".join(["---"] * len(block["headers"])) + " |"])
                lines.extend("| " + " | ".join(row) + " |" for row in block["rows"])
                lines.append("")
    (REPORT / "报告正文.md").write_text("\n".join(lines), encoding="utf-8")


def word(title, subtitle, pages, references, team_name):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = title
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.4)
    section.top_margin = section.bottom_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Pt(21)
    for style_name in ["Heading 1", "Caption"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(12 if style_name == "Heading 1" else 10.5)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = style.paragraph_format.space_after = Pt(0)
    def add(text, style=None, center=False):
        p = doc.add_paragraph(text, style=style)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Pt(0)
        return p
    for _ in range(3):
        add("")
    add("2026 FinTechathon 深圳国际金融科技大赛", center=True)
    add("数据分析赛道 · 赛题三", center=True)
    for _ in range(3):
        add("")
    p = add(title, center=True); p.runs[0].font.size = Pt(20); p.runs[0].bold = True
    add(subtitle, center=True)
    for _ in range(3):
        add("")
    add("小组名称：" + team_name, center=True)
    doc.add_page_break()
    for page_number, blocks in enumerate([*pages, references]):
        if page_number:
            doc.add_page_break()
        for block in blocks:
            if block["type"] == "heading":
                add(block["text"], "Heading 1")
            elif block["type"] == "paragraph":
                p = add(block["text"]); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            elif block["type"] == "figure":
                p = add("", center=True)
                p.add_run().add_picture(str(REPORT / "figures" / block["file"]), width=Cm(16.2 * block["width"]))
                add(block["caption"], "Caption", center=True)
            else:
                add(block["title"], "Caption", center=True)
                t = doc.add_table(rows=1, cols=len(block["headers"])); t.style = "Table Grid"
                for cell, text in zip(t.rows[0].cells, block["headers"]):
                    cell.text = text
                for row in block["rows"]:
                    for cell, text in zip(t.add_row().cells, row):
                        cell.text = text
                for row_index, row in enumerate(t.rows):
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            p.paragraph_format.first_line_indent = Pt(0)
                            p.paragraph_format.line_spacing = 1.5
                            for run in p.runs:
                                run.font.size = Pt(10.5)
                                run.bold = row_index == 0
    doc.save(REPORT / "报告可编辑稿.docx")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--team-name", default="待填写")
    parser.add_argument("--group-id", default="待填编号")
    args = parser.parse_args()
    REPORT.mkdir(exist_ok=True)
    font = ensure_fonts()
    figures(font)
    title, subtitle, pages, references = content()
    markdown(title, subtitle, pages, references, args.team_name)
    word(title, subtitle, pages, references, args.team_name)
    write_tex(title, subtitle, pages, references, args.team_name)
    build = REPORT / "build"
    build.mkdir(exist_ok=True)
    for _ in range(2):
        process = subprocess.run(["xelatex", "-interaction=nonstopmode", "-halt-on-error", "-output-directory=build", "report.tex"], cwd=REPORT, capture_output=True, text=True)
        (build / "compile.log").write_text(process.stdout + process.stderr)
        if process.returncode:
            print(process.stdout[-7000:]); raise RuntimeError("XeLaTeX failed; see report/build/compile.log")
    import shutil
    destination = REPORT / f"{args.group_id}-赛题3.pdf"
    shutil.copy2(build / "report.pdf", destination)
    manifest = {
        "pdf": destination.name, "team_name": args.team_name, "group_id": args.group_id,
        "body_font": "FandolSong-Regular (open-source Song typeface; not Microsoft SimSun)",
        "body_font_size_pt": 10.5, "body_baseline_pt": 15.75,
        "paragraph_indent_pt": 21, "paragraph_spacing_pt": 0,
        "declared_body_pages": 8, "figures": 5,
        "remaining": ["fill actual team name and group ID", "confirm current official data-use terms"],
    }
    (REPORT / "版式与待办说明.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print("Generated", destination)


if __name__ == "__main__":
    main()
