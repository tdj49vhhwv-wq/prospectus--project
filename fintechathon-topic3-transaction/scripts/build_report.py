"""Build the chronological real-order report and five Chinese figures."""
from pathlib import Path
import json, re, subprocess, shutil
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
import numpy as np
from matplotlib.font_manager import FontProperties
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/"outputs"
REPORT=ROOT/"report"

def paragraph(text):
    return {"type": "paragraph", "text": text.replace("−", "-")}


def heading(text):
    return {"type": "heading", "text": text}


def table(title, headers, rows, widths=None):
    return {"type": "table", "title": title, "headers": headers, "rows": rows, "widths": widths}


def figure(name, caption, width=1):
    return {"type": "figure", "file": name, "caption": caption, "width": width}


def escape(text):
    replacements = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
    return "".join(replacements.get(character, character) for character in text)


def tex_block(block):
    kind = block["type"]
    if kind == "heading":
        return r"\section*{" + escape(block["text"]) + "}\n"
    if kind == "paragraph":
        parts = re.split(r"(https?://[A-Za-z0-9./_-]+)", block["text"])
        return "".join(r"\url{" + part.rstrip(".") + "}" + ("." if part.endswith(".") else "") if part.startswith("https://") else escape(part) for part in parts) + "\n\n"
    if kind == "figure":
        return "\\begin{figure}[H]\n\\centering\n\\includegraphics[width=" + str(block["width"]) + r"\linewidth]{figures/" + block["file"].replace(".png", ".pdf") + "}\n\\caption{" + escape(block["caption"]) + "}\n\\end{figure}\n"
    if kind == "table":
        widths = block["widths"] or [1 / len(block["headers"])] * len(block["headers"])
        specification = "".join(r">{\raggedright\arraybackslash}p{" + f"{width:.3f}" + r"\tablewidth}" for width in widths)
        lines = ["\\begin{table}[H]", "\\centering", "\\caption{" + escape(block["title"]) + "}", "\\setlength{\\tablewidth}{\\linewidth}", f"\\addtolength{{\\tablewidth}}{{-{2 * len(widths)}\\tabcolsep}}", "\\begin{tabular}{" + specification + "}", r"\toprule"]
        lines.append(" & ".join(r"\textbf{" + escape(str(x)) + "}" for x in block["headers"]) + r" \\")
        lines.append(r"\midrule")
        lines.extend(" & ".join(escape(str(x)) for x in row) + r" \\" for row in block["rows"])
        lines.extend([r"\bottomrule", r"\end{tabular}", r"\end{table}"])
        return "\n".join(lines) + "\n"
    raise ValueError(kind)


def write_tex(title, subtitle, pages, references, team_name):
    preamble = r"""\documentclass[a4paper,10pt]{article}
\usepackage{fontspec}
\setmainfont[Path=fonts/,BoldFont=FandolSong-Bold.otf,ItalicFont=FandolSong-Regular.otf]{FandolSong-Regular.otf}
\XeTeXlinebreaklocale "zh"
\XeTeXlinebreakskip=0pt plus 1pt
\usepackage[left=24mm,right=24mm,top=22mm,bottom=22mm]{geometry}
\usepackage{graphicx,float,booktabs,array,caption,titlesec,fancyhdr,indentfirst}
\usepackage[hidelinks,unicode]{hyperref}
\hypersetup{pdfauthor={},pdftitle={订单流水与公开评价融合的小微商户经营收缩识别},pdfsubject={赛题三数据分析报告}}
\renewcommand{\normalsize}{\fontsize{10.5bp}{15.75bp}\selectfont}
\normalsize
\setlength{\parindent}{21bp}
\setlength{\parskip}{0pt}
\setlength{\tabcolsep}{4pt}
\newlength{\tablewidth}
\renewcommand{\arraystretch}{1.05}
\setlength{\textfloatsep}{6pt}
\setlength{\intextsep}{6pt}
\captionsetup{labelformat=empty,font=normalsize,skip=3pt}
\titleformat{\section}{\normalfont\fontsize{12pt}{18pt}\selectfont\bfseries}{}{0pt}{}
\titlespacing*{\section}{0pt}{6pt}{4pt}
\pagestyle{fancy}
\fancyhf{}
\fancyfoot[C]{\fontsize{9pt}{12pt}\selectfont\thepage}
\renewcommand{\headrulewidth}{0pt}
\setlength{\headheight}{0pt}
\setlength{\footskip}{12mm}
\emergencystretch=1em
\begin{document}
\begin{titlepage}
\centering
\vspace*{30mm}
{\fontsize{15pt}{23pt}\selectfont 2026 FinTechathon 深圳国际金融科技大赛\par}
\vspace{5mm}
{\fontsize{12pt}{18pt}\selectfont 数据分析赛道 · 赛题三\par}
\vspace{26mm}
"""
    cover = "{\\fontsize{20pt}{30pt}\\selectfont\\bfseries " + escape(title) + "\\par}\n\\vspace{5mm}\n{\\fontsize{13pt}{21pt}\\selectfont " + escape(subtitle) + "\\par}\n\\vspace{30mm}\n{\\fontsize{14pt}{21pt}\\selectfont 小组名称：" + escape(team_name) + "\\par}\n\\end{titlepage}\n\\setcounter{page}{1}\n\\normalsize\n"
    body = "\n\\newpage\n".join("".join(tex_block(block) for block in page) for page in pages)
    tail = "\n\\newpage\n\\fancyfoot[C]{\\fontsize{9pt}{12pt}\\selectfont 参考文献}\n" + "".join(tex_block(block) for block in references) + "\n\\end{document}\n"
    (REPORT / "report.tex").write_text(preamble + cover + body + tail, encoding="utf-8")


def markdown(title, subtitle, pages, references, team_name):
    lines = ["# " + title, subtitle, "", "小组名称：" + team_name, "", "<!-- 封面后正文共8页；参考文献另页；PDF页数已经单独核验。 -->", ""]
    for page_number, blocks in enumerate([*pages, references], 1):
        lines.append(f"<!-- {'正文第'+str(page_number)+'页' if page_number <= 8 else '参考文献'} -->")
        for block in blocks:
            if block["type"] == "heading":
                lines.extend(["## " + block["text"], ""])
            elif block["type"] == "paragraph":
                lines.extend([block["text"], ""])
            elif block["type"] == "figure":
                lines.extend([f"![{block['caption']}](figures/{block['file']})", "", block["caption"], ""])
            else:
                lines.extend([block["title"], "", "| " + " | ".join(block["headers"]) + " |", "| " + " | ".join(["---"] * len(block["headers"])) + " |"])
                lines.extend("| " + " | ".join(row) + " |" for row in block["rows"])
                lines.append("")
    (REPORT / "报告正文.md").write_text("\n".join(lines), encoding="utf-8")


def word(title, subtitle, pages, references, team_name):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc = Document()
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.title = title
    section = doc.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.4)
    section.top_margin = section.bottom_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"; normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Pt(21)
    for style_name in ["Heading 1", "Caption"]:
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(12 if style_name == "Heading 1" else 10.5)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = style.paragraph_format.space_after = Pt(0)
    def add(text, style=None, center=False):
        p = doc.add_paragraph(text, style=style)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.first_line_indent = Pt(0)
        return p
    for _ in range(3):
        add("")
    add("2026 FinTechathon 深圳国际金融科技大赛", center=True)
    add("数据分析赛道 · 赛题三", center=True)
    for _ in range(3):
        add("")
    p = add(title, center=True); p.runs[0].font.size = Pt(20); p.runs[0].bold = True
    add(subtitle, center=True)
    for _ in range(3):
        add("")
    add("小组名称：" + team_name, center=True)
    doc.add_page_break()
    for page_number, blocks in enumerate([*pages, references]):
        if page_number:
            doc.add_page_break()
        for block in blocks:
            if block["type"] == "heading":
                add(block["text"], "Heading 1")
            elif block["type"] == "paragraph":
                p = add(block["text"]); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            elif block["type"] == "figure":
                p = add("", center=True)
                p.add_run().add_picture(str(REPORT / "figures" / block["file"]), width=Cm(16.2 * block["width"]))
                add(block["caption"], "Caption", center=True)
            else:
                add(block["title"], "Caption", center=True)
                t = doc.add_table(rows=1, cols=len(block["headers"])); t.style = "Table Grid"
                for cell, text in zip(t.rows[0].cells, block["headers"]):
                    cell.text = text
                for row in block["rows"]:
                    for cell, text in zip(t.add_row().cells, row):
                        cell.text = text
                for row_index, row in enumerate(t.rows):
                    for cell in row.cells:
                        for p in cell.paragraphs:
                            p.paragraph_format.first_line_indent = Pt(0)
                            p.paragraph_format.line_spacing = 1.5
                            for run in p.runs:
                                run.font.size = Pt(10.5)
                                run.bold = row_index == 0
    doc.save(REPORT / "报告可编辑稿.docx")



def figures(font):
 import matplotlib.dates as md
 folder=REPORT/'figures';folder.mkdir(exist_ok=True)
 s=json.loads((OUT/'analysis_summary.json').read_text());audit=json.loads((OUT/'data_audit.json').read_text());m=pd.read_csv(OUT/'model_metrics.csv',index_col='model');d=pd.read_csv(OUT/'descriptive_stats.csv',index_col='feature');later=pd.read_csv(OUT/'later_anchor_check.csv',index_col='model');ln=json.loads((OUT/'later_anchor_notes.json').read_text())
 plt.rcParams.update({'font.size':9,'axes.unicode_minus':False,'axes.spines.top':False,'axes.spines.right':False,'savefig.dpi':300})
 blue='#296A88';red='#B9573C';green='#287E62';purple='#724D99'
 def finish(fig,name):
  for ax in fig.axes:
   for t in [ax.title,ax.xaxis.label,ax.yaxis.label,*ax.get_xticklabels(),*ax.get_yticklabels(),*ax.texts]:t.set_fontproperties(font)
   if ax.get_legend():
    for t in ax.get_legend().get_texts():t.set_fontproperties(font)
  fig.tight_layout(pad=.8)
  fig.savefig(folder/(name+'.png'),bbox_inches='tight',facecolor='white');fig.savefig(folder/(name+'.pdf'),bbox_inches='tight',facecolor='white');plt.close(fig)
 fig,ax=plt.subplots(figsize=(7.05,2.6));anchors=['2017-06-30','2017-09-30','2017-12-31','2018-03-31','2018-04-30'];numbers=[262,351,457,581,int(later.iloc[0]['n'])]
 for i,date in enumerate(anchors):
  t=pd.Timestamp(date)+pd.Timedelta(days=1);start=t-pd.Timedelta(days=180)
  ax.barh(i,180,left=md.date2num(start),color=blue,height=.3,label='可用历史180天' if i==0 else None)
  ax.barh(i,90,left=md.date2num(t),color=red,height=.3,label='未来90天标签' if i==0 else None);ax.scatter(md.date2num(t),i,c='#333333',s=15)
 ax.set_yticks(range(5),[f'{label}（{n}条）' for label,n in zip(['训练一','训练二','验证','主测试','更晚核对'],numbers)]);ax.invert_yaxis();ax.xaxis_date();ax.xaxis.set_major_locator(md.MonthLocator(interval=3));ax.xaxis.set_major_formatter(md.DateFormatter('%Y-%m'));ax.legend(loc='upper left',bbox_to_anchor=(0,1.21),ncol=2,fontsize=8);ax.set_xlabel('商户—观察时点；更晚核对与主测试窗口有重叠');finish(fig,'figure01_timeline')
 fig,axes=plt.subplots(1,4,figsize=(7.05,2.2))
 for ax,c,title,pct in zip(axes,['days_since_order','review_low_share','known_late_share','undelivered_share'],['距最后订单天数','低分评价占比','已知延期占比','未确认完成占比'],[False,True,True,True]):
  row=d.loc[c];v=np.array([row.normal_mean,row.contraction_mean])*(100 if pct else 1);ax.bar([0,1],v,color=[blue,red],width=.6);ax.set_xticks([0,1],['未达阈值','收缩']);ax.set_title(title,fontsize=9);ax.set_ylim(0,max(v)*1.35)
  for i,x in enumerate(v):ax.text(i,x+max(v)*.04,f'{x:.1f}'+('%' if pct else ''),ha='center',fontsize=8)
 finish(fig,'figure02_profiles')
 fig,axes=plt.subplots(1,2,figsize=(7.05,2.9),sharey=True)
 for ax,metric,title in zip(axes,['roc_auc','ap'],['测试ROC-AUC','测试平均精确率AP']):
  for i,name in enumerate(m.index):
   v=m.loc[name,metric];lo,hi=s['confidence_intervals'][name][metric];c=purple if name==s['primary'] else (blue if name==s['baseline'] else green)
   ax.errorbar(v,i,xerr=[[v-lo],[hi-v]],fmt='o',color=c,capsize=3,markersize=4);ax.text(.97,i,f'{v:.3f}',va='center',ha='right',fontsize=8)
  ax.set_xlim(0,1);ax.set_yticks(range(6),[n.replace('_','·') for n in m.index]);ax.set_title(title);ax.axvline(.5 if metric=='roc_auc' else s['test_prevalence'],c='#888888',ls='--',lw=1);ax.grid(axis='x',alpha=.15)
 axes[0].invert_yaxis();finish(fig,'figure03_models')
 budget=pd.read_csv(OUT/'review_budget.csv',index_col='strategy');pair=json.loads((OUT/'low_rating_comparison.json').read_text());fig,axes=plt.subplots(1,2,figsize=(7.05,2.4),gridspec_kw={'width_ratios':[1.1,1]})
 keys=['random_expected','low_rating',s['baseline'],s['primary']];labels=['随机期望','低评分','交易基线','交易＋公开信息'];values=[budget.loc[n,'captured'] for n in keys]
 axes[0].barh(range(4),values,color=['#AAAAAA','#B88637',blue,purple],height=.6);axes[0].set_yticks(range(4),labels);axes[0].invert_yaxis();axes[0].set_xlim(0,max(values)*1.3);axes[0].set_xlabel(f'复核{int(budget.iloc[0].reviewed)}家时命中数')
 for i,v in enumerate(values):axes[0].text(v+1,i,f'{v:.2f}' if i==0 else str(int(v)),va='center',fontsize=9)
 contrasts=[(pair['primary_minus_low_rating_precision'],pair['paired_ci95']),(m.loc[s['primary'],'precision']-m.loc[s['baseline'],'precision'],s['primary_minus_baseline']['precision'])]
 for i,(v,(lo,hi)) in enumerate(contrasts):axes[1].errorbar(v*100,i,xerr=[[(v-lo)*100],[(hi-v)*100]],fmt='o',color=[green,purple][i],capsize=3)
 axes[1].set_yticks(range(2),['相对低评分','相对交易基线']);axes[1].invert_yaxis();axes[1].axvline(0,c='#777777',ls='--');axes[1].set_xlabel('命中率差（百分点）及95%区间');axes[1].set_xlim(-15,50);finish(fig,'figure04_budget')
 fig,axes=plt.subplots(1,2,figsize=(7.05,2.3))
 for name,label,c in [(s['baseline'],'交易基线',blue),(s['primary'],'增强组合',purple)]:
  v=[m.loc[name,'ap'],later.loc[name,'ap']];axes[0].plot([0,1],v,marker='o',color=c,label=label)
  for i,x in enumerate(v):axes[0].text(i,x+.003,f'{x:.3f}',ha='center',fontsize=8)
 axes[0].set_xticks([0,1],['2018年3月末','2018年4月末']);axes[0].set_xlim(-.3,1.3);axes[0].set_ylim(.48,.62);axes[0].set_ylabel('固定模型AP');axes[0].legend(fontsize=8,loc='lower left')
 for i,(v,ci) in enumerate([(m.loc[s['primary'],'ap']-m.loc[s['baseline'],'ap'],s['primary_minus_baseline']['ap']),(later.loc[s['primary'],'ap']-later.loc[s['baseline'],'ap'],ln['primary_minus_baseline_ap_ci95'])]):
  lo,hi=ci;axes[1].errorbar(v,i,xerr=[[v-lo],[hi-v]],fmt='o',color=purple,capsize=3)
 axes[1].set_yticks([0,1],['主测试','更晚核对']);axes[1].invert_yaxis();axes[1].axvline(0,c='#888888',ls='--');axes[1].set_xlim(-.025,.035);axes[1].set_xlabel('增强相对交易基线的AP差及95%区间');finish(fig,'figure05_stability')


def content():
 s=json.loads((OUT/'analysis_summary.json').read_text());a=json.loads((OUT/'data_audit.json').read_text());m=pd.read_csv(OUT/'model_metrics.csv',index_col='model');b=pd.read_csv(OUT/'review_budget.csv',index_col='strategy');l=pd.read_csv(OUT/'later_anchor_check.csv',index_col='model');ln=json.loads((OUT/'later_anchor_notes.json').read_text());q=json.loads((OUT/'low_rating_comparison.json').read_text());desc=pd.read_csv(OUT/'descriptive_stats.csv',index_col='feature')
 base=s['baseline'];enh=s['primary'];delta=m.loc[enh,'ap']-m.loc[base,'ap'];lo,hi=s['primary_minus_baseline']['ap'];n=s['test_n'];positives=s['test_positive'];k=int(b.loc[enh,'reviewed']);hits=int(b.loc[enh,'captured']);lowhits=int(b.loc['low_rating','captured']);basehits=int(b.loc[base,'captured']);coverage=n/a['cohort_counts'][3]['active_sellers']
 ci=lambda x:f'[{x[0]:.3f}，{x[1]:.3f}]';fmt=lambda v:f'{v:.3f}'
 title='订单流水与公开评价融合的小微商户经营收缩识别';subtitle='——基于真实交易的时间外验证与复核决策'
 pages=[
 [heading('摘要'),paragraph(f'针对小微主体经营数据稀疏的问题，本文采用数据提供方公开的真实电商交易样本，以订单流水构建基线，加入评价及履约信息，检验增量和有限预算复核价值。原始数据含99,441笔订单、112,650条商品明细及99,224条评价；构建17项指标，按季度观察点训练、验证并留出未来90天标签。主测试包含{n}家商户，{positives}家未来平台订单销售额下降至少50%。增强组合AP为{m.loc[enh,"ap"]:.3f}，较交易基线增加{delta:.4f}，95%配对区间为{ci([lo,hi])}，尚未证明稳定增益。复核{k}家时交易基线和增强组合均命中{hits}家，低评分策略命中{lowhits}家；更晚观察点也未确认增强优势。结果支持交易信息用于经营收缩调查排序，公开评价先用于核查解释，不宜强行加分。订单金额不等于银行净现金流，未来平台收缩不是停业或信用违约；本研究属于历史时间外回测。'),paragraph('关键词：公开数据；订单流水；未来经营收缩；时间外验证；复核决策'),heading('一、业务背景与研究问题'),paragraph('小微主体通常缺少标准化报表，交易记录能反映经营规模和波动，公开口碑则可能补充消费体验与履约线索。但评价是否有价值，应在已有交易基线上检验，而不是把差评直接等同风险。本文研究：交易指标能否识别未来平台经营收缩？评价与履约是否增加识别价值？固定核查预算下是否值得增加数据成本？'),paragraph('本研究将电商商户作为信息稀疏经营主体的代理，不声称已确认其小微规模。以未来90天订单销售额收缩作为可验证的经营风险线索，使用真实交易而非模拟流水。公开信息与交易来自同一数据提供方，是模块增量检验，尚非独立外部平台与银行流水的融合验证。')],
 [heading('二、公开数据、清洗与样本'),paragraph('数据来自提供方官方公开仓库[1]，固定提交及文件SHA256见复现清单。README说明这些关系表直接抽样自交易系统；仓库附MIT许可证[2]，报告保留来源，原始消费者记录和逐商户数据不公开。商户、订单、商品明细和评价通过ID关联，不做跨来源猜测匹配。'),table('表1 数据字段、规模与处理规则',['数据表','原始记录数','字段与处理'],[['订单','99,441','批准、配送和预计送达时间；不输入当前最终状态'],['商品明细','112,650','商户ID、订单ID、价格；按商户—订单汇总价格'],['评价','99,224','评分、回答时间；只用观察日前已回答评价'],['商户','3,095','匿名商户ID、地区；地区仅作敏感性核对']],[.18,.18,.64]),paragraph(f'订单缺失批准时间160笔，无法定位交易时点，排除其关联交易；商品明细均匹配订单。复核发现61笔配送早于批准，将配送时间记为未知；20条评价早于批准，剔除。1,278笔多商户订单的评价不能可靠归因，评价特征只用单商户订单，交易仍按实际商户明细分配。批准交易时间覆盖2016年9月至2018年9月。'),figure('figure01_timeline.png','图1 历史可用窗口与未来标签窗口；更晚核对并非独立复验'),paragraph(f'仅按观察日前条件纳入：过去180天订单≥10笔、近90天≥3笔、首笔交易至少早于观察日90天。主测试期近180天有交易商户1,686家，纳入{n}家，覆盖{coverage:.2%}；低频、新进入或信息不足商户另行补资料，不能默认为低风险。')],
 [heading('三、指标机制与描述性分析'),paragraph('交易指标T衡量经营规模、变化和稳定性；评价Q反映消费者体验；履约L提供交付质量和积压线索。机制假设是体验与交付问题可能影响复购和经营，但这种关系受规模、季节性和平台结构影响，必须验证而不硬编码风险方向。'),table('表2 交易、评价与履约指标体系（17项）',['模块','指标','口径与机制'],[['T：交易9项','180/90天金额与订单数；金额/订单变化；月度金额变异系数；客单价；距最后订单天数','批准订单商品价之和；变化用log(1+近期)-log(1+前期)，衡量经营动态'],['Q：评价5项','均分、低分占比、评分标准差、近期均分变化、评价覆盖率','仅观察日前已回答的单商户订单评价；低分≤2分，反映体验线索'],['L：履约3项','已知延期占比、已知配送天数、未确认完成占比','只按当时已发生的配送时间判断；后来配送不提前视为已完成']],[.17,.43,.40]),paragraph('金额不含运费，不等于客户总支付、商户结算收入或净现金流，退款和取消历史不可得。评价使用回答时间而非问卷创建时间；同订单多次评价只取当时最新回答。评分变化按订单所属近期/前期窗口计算；无评价或配送记录时保留缺失，仅在拟合数据中估计中位数并增加缺失指示。'),figure('figure02_profiles.png','图2 主测试期商户的历史指标均值（按未来收缩标签分组）'),paragraph(f'未来收缩组距最后订单均值为{desc.loc["days_since_order","contraction_mean"]:.2f}天，未达阈值组{desc.loc["days_since_order","normal_mean"]:.2f}天；低分占比分别为{desc.loc["review_low_share","contraction_mean"]:.2%}和{desc.loc["review_low_share","normal_mean"]:.2%}。未确认完成比例的方向并不符合简单“越高越危险”的假设。以上仅作描述，不证明因果或独立增量。')],
 [heading('四、未来标签与时间外验证'),paragraph('统一月末观察，特征只用观察日及之前180天可获得事件。未来90天商品销售额记为G未来，过去90天记为G历史，主标签为G未来≤0.5×G历史。标签代表平台内显著经营收缩，包含可能的平台迁移或正常波动；不称为真实财务困境、停业或违约。门槛在建模前固定，并用下降≥70%与≥30%定义作敏感性分析。'),table('表3 时间划分、样本与标签成熟',['阶段／观察日','商户—时点数','收缩数','标签区间'],[['训练：2017-06-30','262','75','2017-07-01至09-28'],['训练：2017-09-30','351','86','2017-10-01至12-29'],['验证：2017-12-31','457','135','2018-01-01至03-31'],['主测试：2018-03-31',str(n),str(positives),'2018-04-01至06-29'],['更晚核对：2018-04-30',str(int(l.iloc[0]['n'])),str(int(l.iloc[0]['positive'])),'2018-05-01至07-29']],[.31,.18,.12,.39]),paragraph('训练613条商户—时点记录来自389家商户；验证457家。主测试一商户一条记录，训练标签在验证前成熟，验证标签在主测试观察日结束时成熟。相同商户可出现在不同阶段，检验的是以后时点而非随机商户划分；更晚核对与主测试商户及未来窗口重叠，不冒充独立复验。'),paragraph('比较LR逻辑回归与HGB梯度提升树，各含T、TQ、TQL三种输入。LR调C=0.1、1、10；HGB调叶节点7、15和L2=1、5，固定150次迭代、学习率0.05、每叶至少20条，不早停。只按验证AP调参，并按T基线验证AP选算法家族，再固定该家族比较公开信息增量。本次选择LR，TQL是主增量检验组合，并非验证最优部署方案。'),paragraph('最终模型用训练＋验证重拟合，预处理只在拟合数据学习；不用主测试标签调参。报告ROC-AUC、非插值AP、Brier及约10%预算指标；2,000次配对商户Bootstrap覆盖固定模型的本次测试抽样不确定性，不涵盖训练重拟合和选择过程。更晚核对沿用同一固定模型，未加入主测试标签重训。')],
 [heading('五、交易基线与公开信息的增量'),figure('figure03_models.png','图3 六个预设模型的测试表现及95%配对Bootstrap区间'),table('表4 验证及主测试结果（Brier越低越好）',['模型','验证AP','测试AUC','测试AP','Brier'],[[name.replace('_','·'),fmt(row.validation_ap),fmt(row.roc_auc),fmt(row.ap),f'{row.brier:.4f}'] for name,row in m.iterrows()],[.26,.17,.19,.18,.20]),paragraph(f'交易基线{base.replace("_","·")}的测试AP为{m.loc[base,"ap"]:.3f}，增强组合{enh.replace("_","·")}为{m.loc[enh,"ap"]:.3f}；AP增量{delta:.4f}，95%区间{ci([lo,hi])}包含0。AUC从{m.loc[base,"roc_auc"]:.3f}到{m.loc[enh,"roc_auc"]:.3f}，差的区间也包含0。评价Q和履约L的点估计改善不能证明已提供稳定增益，未据测试最高分回头改选模型。'),paragraph(f'验证期T基线AP高于TQL增强，部署选择不能忽略这一结果。拟合期正例比例{s["fitting_prevalence"]:.2%}，主测试{s["test_prevalence"]:.2%}，存在分布变化。用拟合期比例作常数预测的测试Brier为{s["constant_brier"]:.4f}，增强组合为{m.loc[enh,"brier"]:.4f}；这属于概率评分比较，不证明概率已校准，更不能转称信用违约概率。')],
 [heading('六、应用验证：有限预算复核'),paragraph(f'预先设定每期约10%核查预算，{n}家向上取整复核{k}家（{k/n:.2%}）。名单按固定分数排序，不用未来标签选名单；低评分策略使用同一历史评价口径，并按预设匿名ID顺序处理并列。未来标签只用于事后核对调查优先级。'),figure('figure04_budget.png','图4 固定预算命中数及增强组合相对两类基线的命中率差'),table('表5 主测试复核预算效果',['策略','命中数','命中率','捕获率','随机提升'],[[label,f'{b.loc[name,"captured"]:.2f}' if name=='random_expected' else str(int(b.loc[name,'captured'])),f'{b.loc[name,"precision"]:.2%}',f'{b.loc[name,"recall"]:.2%}',f'{b.loc[name,"lift"]:.2f}'] for name,label in [('random_expected','随机期望'),('low_rating','低评分'),(base,'交易基线'),(enh,'交易＋公开信息')]],[.30,.16,.19,.19,.16]),paragraph(f'增强组合相对低评分多命中{hits-lowhits}家，命中率差{q["primary_minus_low_rating_precision"]*100:.2f}个百分点，配对区间为[{q["paired_ci95"][0]*100:.2f}，{q["paired_ci95"][1]*100:.2f}]个百分点。但交易基线同样命中{basehits}家，增加公开信息的复核命中数增量为{hits-basehits}，不能把交易优势归给评价模块。随机{b.loc["random_expected","captured"]:.2f}家为理论期望。'),paragraph(f'增强名单仍漏掉{1-b.loc[enh,"recall"]:.2%}的未来收缩商户，因此不能用于排除未上榜主体的风险。这里只验证未来平台收缩线索的排序，不测真实违约、财务损失或处置收益；复核命中不能直接换算为节省金额。')],
 [heading('七、稳定性复核与适用边界'),figure('figure05_stability.png','图5 固定模型跨观察点表现；两个窗口重叠，不能视为独立复验'),paragraph(f'更晚核对含{int(l.iloc[0]["n"])}家商户，交易基线AP为{l.loc[base,"ap"]:.3f}，增强组合{l.loc[enh,"ap"]:.3f}，增量区间{ci(ln["primary_minus_baseline_ap_ci95"])}仍包含0。两期增强增量均小，尚未证明稳定改善；更晚期AUC增强未优于交易基线。反复选择最有利月份或标签门槛会产生选择偏差，本研究保留所有核对结果。'),paragraph('按订单≥20、评价覆盖≥50%、主要地区/其他地区及训练未出现商户进行固定预测核对；标签改为下降≥70%时，增强AP没有优于基线，改为下降≥30%时点估计增加。子组没有重新调参，也没有据此改主结论；阈值敏感性与地区核对不等于独立外推验证。'),paragraph('数据边界：数据提供方只发布平台抽样，不是全部经营额；未来无订单可能表示平台迁移、样本缺失或主动退出。价格及订单明细缺少事件版本，退款、取消历史不可得，无法完全还原历史结算状态。按时间戳重建可用信息比静态营业标签更严格，但仍是历史回测。'),paragraph('外推边界：商户规模未确认，海外2016—2018年电商数据不能证明2026年深圳线下小微主体适用。评价和履约同源，不能据此证明跨独立平台公开数据的增益。低频商户和新商户被门槛排除，应单独核查；全量漏报须有独立标签或带已知抽样概率的随机核查，不能只核查模型名单后声称召回率提高。')],
 [heading('八、产品思路与商业闭环'),paragraph('建议产品定位为“商户经营线索助手”：用交易变化形成待调查名单，用公开评价及履约线索提示核查方向。当前证据优先支持交易基线排序，公开模块先作解释展示；不把未确认增量的模块强行叠加评分，不输出贷款违约概率。'),table('表6 证据对应产品动作与验收',['模块','已有证据','业务动作'],[['可用性分流',f'主测试覆盖{coverage:.2%}活跃商户','未覆盖者补资料并保留随机核查'],['交易调查排序',f'{k}家命中{basehits}家未来收缩','核实收入、退款及平台迁移，独立记录结果'],['评价／履约画像','增量区间包含0；同预算未增加命中','展示真实历史指标及来源，不作为拒贷依据'],['反馈与回退','尚无真实信用表现或处置成本','记录误报、失访、耗时及干预；效果不足暂停排序']],[.23,.34,.43]),paragraph('商业验收设每期复核k家，p为经未来核实的可行动风险命中率，v为每次有效发现的处置净价值，c为单家核查成本，F为固定成本，则净价值=k×(p×v-c)-F。同预算增加公开模块的增量价值=(新增有效命中)×v-新增固定成本。本次增量命中为0，不能证明值得增加评分成本；画像是否节省核查时间需另行试点测量。'),heading('九、结论与下一步验证'),paragraph('第一，真实订单基线与未来90天标签补强了研究契合度与时间方向，交易排序优于单纯低评分筛选。第二，公开评价及履约的增量仍未得到稳定证明；合理决策是将其作为核查线索，保留交易基线。第三，未来平台收缩不等于财务困境或违约，应在取得授权的真实结算流水和未来信用标签上复验。'),paragraph('后续在统一月末锁定真实流水基线、流水＋独立公开信息、10%预算及成本口径，训练标签须在测试日前成熟；留出后续月份和新地区，保留随机核查对照。未知/失访不标为正常，干预效果单独评价，每月监测覆盖、命中和成本；仅在预设增量区间下界>0且净价值>0时扩大，否则回退人工核实。该闭环是待执行试点，不是已完成业务成果。')]
 ]
 refs=[heading('参考文献与数据来源'),paragraph('[1] Olist. work-at-olist-data: README及公开交易数据[DS/OL]. https://github.com/olist/work-at-olist-data. 固定提交d9e49802f3e92d09ee94ab9ccc5e457f207a8959；官方说明关系表直接抽样自交易环境，非模拟订单。访问日期2026-10-12。'),paragraph('[2] Olist. MIT License[EB/OL]. Copyright (c) 2019 @olist. https://github.com/olist/work-at-olist-data/blob/d9e49802f3e92d09ee94ab9ccc5e457f207a8959/LICENSE. 原始许可及文件SHA256已随复现材料保存；报告只展示匿名聚合结果，不发布原始消费者记录。'),paragraph('[3] scikit-learn. Model evaluation; LogisticRegression; HistGradientBoostingClassifier[EB/OL]. https://scikit-learn.org/stable/. AP采用average_precision_score，Brier为概率预测均方误差。'),paragraph('[4] Efron B, Tibshirani R J. An Introduction to the Bootstrap[M]. New York: Chapman & Hall, 1993.'),paragraph('[5] Brier G W. Verification of forecasts expressed in terms of probability[J]. Monthly Weather Review, 1950, 78(1): 1–3.')]
 return title,subtitle,pages,refs


def main():
 REPORT.mkdir(exist_ok=True);fontfolder=REPORT/'fonts';fontfolder.mkdir(exist_ok=True)
 import hashlib, requests
 fonts={'FandolSong-Regular.otf':('FandolSong-Regular/FandolSong-Regular.otf','5dc10d4d76ec3bc5fcc6065d80088c38eee0cbf851bee5b5cad1334a796a0d59'),'FandolSong-Bold.otf':('FandolSong-Bold/FandolSong-Bold.otf','ba71e3247d04ffcd547f5826c7eb747da4f70cfe047c166283f5f31c402817eb'),'COPYING':('COPYING','853b586f0d520493390e571431afaf36a5fbb27dcfd239338a7ee9b0505cb004')}
 old=ROOT.parent/'fintechathon-topic3/report/fonts'
 for name,(remote,digest) in fonts.items():
  dest=fontfolder/name
  if not dest.exists():
   if (old/name).exists():shutil.copy2(old/name,dest)
   else:
    response=requests.get('https://raw.githubusercontent.com/Yixf-Self/fandol-fonts/b93300821373d3092e378e57ae22ecb8e9082c62/'+remote,timeout=60);response.raise_for_status();dest.write_bytes(response.content)
  assert hashlib.sha256(dest.read_bytes()).hexdigest()==digest
 font=FontProperties(fname=str(fontfolder/'FandolSong-Regular.otf'));figures(font);title,subtitle,pages,refs=content();markdown(title,subtitle,pages,refs,'285小组');word(title,subtitle,pages,refs,'285小组');write_tex(title,subtitle,pages,refs,'285小组')
 build=REPORT/'build';build.mkdir(exist_ok=True)
 for _ in range(2):
  p=subprocess.run(['xelatex','-interaction=nonstopmode','-halt-on-error','-output-directory=build','report.tex'],cwd=REPORT,capture_output=True,text=True);(build/'compile.log').write_text(p.stdout+p.stderr)
  if p.returncode:raise RuntimeError(p.stdout[-6000:])
 shutil.copy2(build/'report.pdf',REPORT/'285-赛题3.pdf')
 settings={'pdf':'285-赛题3.pdf','team_name':'285小组','group_id':'285','body_font':'FandolSong-Regular','body_pages':8,'figures':5,'remaining':['Confirm required Song font acceptance; re-export with SimSun if required'],'new_scope':'Real ecommerce orders, future 90-day platform-sales contraction; not bank cash flow or credit default'}
 (REPORT/'版式与待办说明.json').write_text(json.dumps(settings,ensure_ascii=False,indent=2));print('Generated',REPORT/'285-赛题3.pdf')

if __name__=='__main__':main()
